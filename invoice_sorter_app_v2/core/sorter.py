import io
import logging
import os
import re
import shutil
import sys
import threading
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path

import fitz  # PyMuPDF
import numpy as np
from PIL import Image


DATE_FOLDER_RE = re.compile(r"^\d{1,2}-[A-Za-z]{3}-\d{2}$", re.I)
OCR_SCALE = 1.2
OCR_RETRY_SCALE = 1.7
HEADER_FRACTION = 0.4
HEADER_TALL_FRACTION = 0.55
_PADDLE_LOCK = threading.Lock()
_PADDLE_ENGINE = None
_RAPID_LOCAL = threading.local()
_OCR_LINES = threading.local()
os.environ.setdefault("PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT", "0")
os.environ.setdefault("FLAGS_use_mkldnn", "0")
os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("MKL_NUM_THREADS", "1")


def app_root() -> Path:
    """Folder that contains the customer mapping (source tree or frozen EXE)."""
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).resolve().parent
        meipass = getattr(sys, "_MEIPASS", None)
        candidates = [exe_dir]
        if meipass:
            candidates.append(Path(meipass))
        for folder in candidates:
            if (folder / "customer_master_alias_mapping.xlsx").is_file() or (folder / "customers.txt").is_file():
                return folder
        return exe_dir
    return Path(__file__).resolve().parents[1]


def worker_count():
    # Pytest stays single-threaded so OCR tests stay stable.
    if os.environ.get("PYTEST_CURRENT_TEST"):
        return 1
    override = os.environ.get("INVOICE_SORTER_WORKERS", "").strip()
    if override.isdigit():
        return max(1, int(override))
    cpu = os.cpu_count() or 2
    return max(1, min(2, cpu))


def _ort_intra_threads() -> int:
    value = os.environ.get("INVOICE_SORTER_ORT_INTRA_THREADS", "").strip()
    return max(1, int(value)) if value.isdigit() else 2


def _ort_inter_threads() -> int:
    value = os.environ.get("INVOICE_SORTER_ORT_INTER_THREADS", "").strip()
    return max(1, int(value)) if value.isdigit() else 1


def debug_ocr_enabled() -> bool:
    return os.environ.get("INVOICE_SORTER_DEBUG_OCR", "").strip().lower() in {"1", "true", "yes"}


def _debug_ocr(message: str, **fields):
    if not debug_ocr_enabled():
        return
    parts = [message]
    for key, value in fields.items():
        if value is None or value == "":
            continue
        parts.append(f"{key}={value}")
    logging.getLogger("invoice_sorter.ocr").info(" | ".join(parts))


def paddle_retry_enabled():
    return os.environ.get("INVOICE_SORTER_USE_PADDLE", "").strip().lower() in {"1", "true", "yes"}


def get_rapid_engine():
    """Fast ONNX OCR. One engine per thread so PDFs can be processed in parallel."""
    engine = getattr(_RAPID_LOCAL, "engine", None)
    if engine is None:
        from rapidocr import RapidOCR

        engine = RapidOCR(
            params={
                "Global.use_cls": False,
                "Global.max_side_len": 960,
                "Global.log_level": "error",
                "EngineConfig.onnxruntime.use_cuda": False,
                "EngineConfig.onnxruntime.intra_op_num_threads": _ort_intra_threads(),
                "EngineConfig.onnxruntime.inter_op_num_threads": _ort_inter_threads(),
            }
        )
        _RAPID_LOCAL.engine = engine
    return engine


def get_paddle_engine():
    """Slower CPU OCR used only for the first-page retry when explicitly enabled."""
    global _PADDLE_ENGINE
    if _PADDLE_ENGINE is False:
        return None
    if _PADDLE_ENGINE is None:
        with _PADDLE_LOCK:
            if _PADDLE_ENGINE is None:
                try:
                    from paddleocr import PaddleOCR

                    _PADDLE_ENGINE = PaddleOCR(
                        lang="en",
                        use_doc_orientation_classify=False,
                        use_doc_unwarping=False,
                        use_textline_orientation=False,
                    )
                except Exception:
                    _PADDLE_ENGINE = False
    return None if _PADDLE_ENGINE is False else _PADDLE_ENGINE


def get_ocr_engine():
    """RapidOCR for the normal path. Kept for tests and fallbacks."""
    return get_rapid_engine()


def _paddle_result_to_text(result) -> str:
    if not result:
        return ""
    first = result[0] if isinstance(result, list) else result
    rec_texts = None
    if hasattr(first, "get"):
        rec_texts = first.get("rec_texts")
    if rec_texts is None:
        rec_texts = getattr(first, "rec_texts", None)
    if rec_texts:
        return "\n".join(str(t) for t in rec_texts)
    lines = result[0] if isinstance(result, list) and result else []
    texts = []
    for item in lines or []:
        if item and len(item) >= 2 and item[1]:
            texts.append(str(item[1][0]))
    return "\n".join(texts)


@dataclass(frozen=True)
class OcrLine:
    text: str
    box: tuple[float, float, float, float]
    score: float | None = None


def _rect_from_box(box) -> tuple[float, float, float, float]:
    arr = np.asarray(box, dtype=float)
    if arr.size == 0:
        return (0.0, 0.0, 0.0, 0.0)
    if arr.ndim == 1 and arr.size == 4:
        return (float(arr[0]), float(arr[1]), float(arr[2]), float(arr[3]))
    pts = arr.reshape(-1, 2)
    return (
        float(pts[:, 0].min()),
        float(pts[:, 1].min()),
        float(pts[:, 0].max()),
        float(pts[:, 1].max()),
    )


def _rapid_output(result):
    """Normalize RapidOCR's object or tuple result to boxes, texts, scores."""
    if result is None:
        return None, None, None
    boxes = getattr(result, "boxes", None)
    txts = getattr(result, "txts", None)
    scores = getattr(result, "scores", None)
    if txts is not None:
        return boxes, txts, scores
    if isinstance(result, (list, tuple)):
        if len(result) >= 3 and not isinstance(result[0], (int, float)):
            return result[0], result[1], result[2]
        if len(result) == 2 and hasattr(result[0], "txts"):
            return _rapid_output(result[0])
    return None, None, None


def lines_from_rapid_result(result) -> list[OcrLine]:
    boxes, txts, scores = _rapid_output(result)
    if not txts:
        return []
    lines = []
    for index, text in enumerate(txts):
        value = str(text).strip()
        if not value:
            continue
        if boxes is not None and index < len(boxes):
            box = _rect_from_box(boxes[index])
        else:
            box = (0.0, float(index) * 20.0, 800.0, float(index) * 20.0 + 16.0)
        score = None
        if scores is not None and index < len(scores) and scores[index] is not None:
            score = float(scores[index])
        lines.append(OcrLine(value, box, score))
    return lines


def ocr_image_rapid_lines(image: Image.Image) -> list[OcrLine]:
    array = np.asarray(image.convert("RGB"))
    result = get_rapid_engine()(array)
    return lines_from_rapid_result(result)


def ocr_image_rapid(image: Image.Image) -> str:
    lines = ocr_image_rapid_lines(image)
    _OCR_LINES.lines = lines
    return "\n".join(line.text for line in lines)


def take_last_ocr_lines() -> list[OcrLine]:
    lines = getattr(_OCR_LINES, "lines", None) or []
    _OCR_LINES.lines = None
    return list(lines)


def ocr_image_paddle(image: Image.Image) -> str:
    engine = get_paddle_engine()
    if engine is None:
        return ""
    array = np.asarray(image.convert("RGB"))
    with _PADDLE_LOCK:
        if hasattr(engine, "predict"):
            result = engine.predict(array)
        else:
            result = engine.ocr(array, cls=False)
        return _paddle_result_to_text(result)


def ocr_image(image: Image.Image) -> str:
    return ocr_image_rapid(image)


def safe_name(value: str) -> str:
    value = re.sub(r"\s+", " ", value or "").strip()
    value = re.sub(r'[<>:"/\\|?*]', "_", value)
    value = value.rstrip(". ")
    return value or "UNKNOWN"


def parse_date_folder(folder_name: str):
    for fmt in ("%d-%b-%y", "%d-%B-%y"):
        try:
            return datetime.strptime(folder_name, fmt).date()
        except ValueError:
            pass
    return None


def year_from_scan_folder(folder_name: str) -> int | None:
    """Output year comes from the source scan-date folder, not the printed invoice date.

    ``01-Sep-26`` → 2026, ``03-Jan-27`` → 2027. ``25 September`` has no year and
    returns None so the document goes to REVIEW.
    """
    parsed = parse_date_folder(folder_name)
    if parsed:
        return parsed.year
    text = (folder_name or "").strip()
    for fmt in ("%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).year
        except ValueError:
            continue
    return None


def render_page_band(page, scale=OCR_SCALE, fraction=HEADER_FRACTION):
    """Render only the invoice header band needed for identification."""
    rect = page.rect
    clip = fitz.Rect(0, 0, rect.width, max(1, rect.height * fraction))
    pix = page.get_pixmap(
        matrix=fitz.Matrix(scale, scale),
        clip=clip,
        alpha=False,
    )
    return Image.frombytes("RGB", [pix.width, pix.height], pix.samples)


def lines_from_pymupdf(page) -> list[OcrLine]:
    """Line boxes from an embedded text layer. Scanned pages return nothing."""
    data = page.get_text("dict") or {}
    lines = []
    for block in data.get("blocks") or []:
        for line in block.get("lines") or []:
            text = "".join(span.get("text", "") for span in line.get("spans") or []).strip()
            if not text:
                continue
            box = line.get("bbox") or (0, 0, 0, 0)
            lines.append(OcrLine(text, (float(box[0]), float(box[1]), float(box[2]), float(box[3]))))
    return lines


def lines_from_plain_text(text: str) -> list[OcrLine]:
    """Stack plain lines vertically when a PDF has text but no boxes."""
    lines = []
    for index, raw in enumerate((text or "").splitlines()):
        value = raw.strip()
        if not value:
            continue
        top = float(index) * 20.0
        lines.append(OcrLine(value, (0.0, top, 800.0, top + 16.0)))
    return lines


def ocr_scanned_page(page, scale: float = OCR_SCALE) -> str:
    """OCR only the GST header band of the supplied page."""
    return ocr_image_rapid(render_page_band(page, scale=scale, fraction=HEADER_FRACTION))


def ocr_first_page(pdf_path: Path, scale: float = OCR_SCALE):
    """Read embedded text or OCR ONLY page 1 of a source PDF.

    Every source PDF is one complete invoice package. The remaining pages are
    supporting documents and are never rendered/OCR'd. The third value is the
    page-1 lines with boxes, used to find the billed-to customer.
    """
    doc = fitz.open(pdf_path)
    if len(doc) == 0:
        doc.close()
        return "", 0, []

    page = doc[0]
    embedded = page.get_text("text").strip()
    if len(embedded) >= 40:
        text = embedded
        lines = lines_from_pymupdf(page) or lines_from_plain_text(text)
    else:
        _OCR_LINES.lines = None
        text = ocr_scanned_page(page, scale=scale)
        lines = take_last_ocr_lines() or lines_from_plain_text(text)
    page_count = len(doc)
    doc.close()
    return text, page_count, lines


def ocr_pdf(pdf_path: Path, scale: float = OCR_SCALE):
    """Backward-compatible name: now reads ONLY the first page."""
    text, _page_count, _lines = ocr_first_page(pdf_path, scale=scale)
    return [(0, text)]


def retry_first_page_read(pdf_path: Path, first_page_text: str):
    """Retry page 1 at a taller clip. Returns text plus OCR lines when scanned."""
    doc = fitz.open(pdf_path)
    if len(doc) == 0:
        doc.close()
        return first_page_text, []

    page = doc[0]
    embedded = page.get_text("text").strip()
    # If real embedded text exists, don't waste time OCR'ing it again.
    if len(embedded) >= 40:
        lines = lines_from_pymupdf(page) or lines_from_plain_text(first_page_text)
        doc.close()
        return first_page_text, lines

    image = render_page_band(
        page,
        scale=OCR_RETRY_SCALE,
        fraction=HEADER_TALL_FRACTION,
    )
    _OCR_LINES.lines = None
    retried = ocr_image_rapid(image)
    lines = take_last_ocr_lines()
    if paddle_retry_enabled() and not resolve_invoice_number(retried, lines):
        paddle_text = ocr_image_paddle(image)
        if paddle_text:
            retried = paddle_text
            lines = lines_from_plain_text(paddle_text)
    doc.close()
    return retried, lines or lines_from_plain_text(retried)


def retry_ocr_first_page(pdf_path: Path, first_page_text: str):
    """Retry OCR at higher resolution, but still ONLY on page 1."""
    text, _lines = retry_first_page_read(pdf_path, first_page_text)
    return text


def retry_ocr_without_invoice_starts(pdf_path: Path, page_texts):
    """Backward-compatible wrapper; retries only the first page."""
    if not page_texts:
        return []
    page_no, text = page_texts[0]
    return [(page_no, retry_ocr_first_page(pdf_path, text))]


def normalize_ocr_text(text: str) -> str:
    text = (text or "").replace("\u00a0", " ")
    text = re.sub(r"(?i)lnvoice", "Invoice", text)
    text = re.sub(r"(?i)InvoiceNo", "Invoice No", text)
    text = re.sub(r"(?i)Inv\.?\s*No", "Invoice No", text)
    return text


_PAN_RE = re.compile(r"^[A-Z]{5}\d{4}[A-Z]$", re.I)
_GSTIN_RE = re.compile(r"^\d{2}[A-Z]{5}\d{4}[A-Z]\dZ[A-Z0-9]$", re.I)
_PRINTED_DATE_RE = re.compile(r"\b(\d{1,2})[./-](\d{1,2})[./-](\d{4}|\d{2})\b")


def _plausible_gst_invoice_number(value: str) -> bool:
    value = (value or "").strip(" .,:;")
    if not value:
        return False
    if _PAN_RE.fullmatch(value) or _GSTIN_RE.fullmatch(value):
        return False
    if re.search(r"PAN|GSTIN", value, re.I):
        return False
    if re.fullmatch(r"20\d{9}", value):
        return True
    if re.fullmatch(r"\d{8,14}", value):
        return True
    return False


def extract_invoice_number(text: str):
    """Extract the printed GST invoice number, not PAN/GSTIN from the same header row."""
    text = normalize_ocr_text(text)
    labeled = re.search(
        r"Invoice\s*No\.?\s*(?:&\s*Date)?[\s:\-]*((?:(?!Invoice).){0,160})",
        text,
        flags=re.I | re.S,
    )
    windows = [labeled.group(1)] if labeled else []
    windows.append(text)

    for window in windows:
        match = re.search(r"\b(20\d{9})\b", window)
        if match:
            return match.group(1)

    patterns = [
        r"Invoice\s*No\.?\s*(?:&\s*Date)?\s*[:\-]?\s*([0-9A-Z][0-9A-Z./_-]{5,})",
        r"Invoice\s*(?:Number|#)\s*[:\-]?\s*([0-9A-Z][0-9A-Z./_-]{5,})",
    ]
    for pattern in patterns:
        for match in re.finditer(pattern, text, flags=re.I):
            invoice_no = match.group(1).strip(" .,:;")
            invoice_no = re.split(r"\s*[-–]\s*\d{1,2}[/-]\d{1,2}", invoice_no, maxsplit=1)[0]
            invoice_no = invoice_no.strip(" .,:;")
            if _plausible_gst_invoice_number(invoice_no):
                return invoice_no
    return None


_INVOICE_LABEL_PREFIX = re.compile(
    r"(?i)^(?:Invoice\s*No\.?|Invoice\s*(?:Number|#)|Inv\.?\s*No\.?)\s*(?:&\s*Date)?\s*[:\-&]*\s*",
)


def _invoice_anchor_text(value: str) -> str:
    text = normalize_ocr_text(value or "").upper()
    text = re.sub(r"[^A-Z0-9#]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _invoice_label_rank(value: str) -> int:
    norm = _invoice_anchor_text(value)
    if not norm:
        return 0
    compact = norm.replace(" ", "")
    if "INVOICENO" in compact or "INVOICENUMBER" in compact or "INVOICE#" in compact:
        return 3
    if "INVNO" in compact or norm.startswith("INV NO"):
        return 2
    if norm.startswith("INVOICE"):
        return 1
    return 0


def _invoice_candidates_in_text(text: str) -> list[str]:
    text = normalize_ocr_text(text or "")
    stripped = _INVOICE_LABEL_PREFIX.sub("", text).strip()
    bodies = [stripped] if stripped != text.strip() else []
    bodies.append(text)
    found: list[str] = []
    for body in bodies:
        for match in re.finditer(r"\b([0-9A-Z][0-9A-Z./_-]{5,})\b", body):
            token = match.group(1).strip(" .,:;")
            token = re.split(r"\s*[-–]\s*\d{1,2}[/-]\d{1,2}", token, maxsplit=1)[0]
            token = token.strip(" .,:;")
            if _plausible_gst_invoice_number(token):
                found.append(token)
        for match in re.finditer(r"\b(\d{8,14})\b", body):
            token = match.group(1)
            if _plausible_gst_invoice_number(token):
                found.append(token)
    # Preserve order, drop duplicates.
    seen: set[str] = set()
    ordered: list[str] = []
    for item in found:
        if item not in seen:
            seen.add(item)
            ordered.append(item)
    return ordered


def extract_invoice_number_from_lines(lines: list[OcrLine]) -> str | None:
    """Spatial pass: invoice label boxes, then number on same row / below / on label line."""
    if not lines:
        return None
    ranked = [(_invoice_label_rank(line.text), index, line) for index, line in enumerate(lines)]
    anchors = [item for item in ranked if item[0] > 0]
    if not anchors:
        return None
    _rank, _index, anchor = max(anchors, key=lambda item: (item[0], -item[2].box[1]))
    ax0, ay0, ax1, ay1 = anchor.box
    reject_reasons: list[str] = []
    candidates: list[tuple[int, float, float, str, str]] = []

    for token in _invoice_candidates_in_text(anchor.text):
        candidates.append((0, ay0, ax0, token, "on_label"))

    for line in lines:
        if line is anchor:
            continue
        x0, y0, x1, y1 = line.box
        for token in _invoice_candidates_in_text(line.text):
            same_row = _same_row(anchor.box, line.box) and x0 >= ax0 - 4
            below = y0 >= ay1 - 2 and _x_overlap(anchor.box, line.box) >= 0.2
            right = same_row and x0 >= ax1 - 12
            if same_row and right:
                relation = "right"
                priority = 1
            elif below:
                relation = "below"
                priority = 2
            elif same_row:
                relation = "same_row"
                priority = 3
            else:
                reject_reasons.append(f"spatial_reject:{token}:no_relation")
                continue
            score = line.score if line.score is not None else 0.0
            candidates.append((priority, y0, -score, token, relation))

    if not candidates:
        _debug_ocr(
            "invoice_spatial_miss",
            labels=[line.text for _, _, line in anchors[:3]],
            reject=";".join(reject_reasons[:8]),
        )
        return None

    candidates.sort(key=lambda item: (item[0], item[1], item[2]))
    chosen = candidates[0][3]
    _debug_ocr(
        "invoice_spatial_hit",
        value=chosen,
        relation=candidates[0][4],
        label=anchor.text,
    )
    return chosen


def resolve_invoice_number(
    text: str,
    lines: list[OcrLine] | None = None,
    *,
    source: str = "page1",
    pdf_path: Path | None = None,
) -> str | None:
    """Pass 1 regex, then pass 2 spatial. Debug logs when both fail."""
    lines = lines or []
    invoice_no = extract_invoice_number(text)
    if invoice_no:
        return invoice_no
    spatial = extract_invoice_number_from_lines(lines)
    if spatial:
        return spatial
    if debug_ocr_enabled():
        label_lines = [line.text for line in lines if _invoice_label_rank(line.text) > 0]
        sample = normalize_ocr_text(text or "")[:400].replace("\n", " ")
        _debug_ocr(
            "invoice_extract_failed",
            source=source,
            pdf=pdf_path.name if pdf_path else "",
            text_sample=sample,
            label_lines=" | ".join(label_lines[:5]),
            line_boxes=len(lines),
            regex_candidates=_invoice_candidates_in_text(text),
            spatial_attempted=bool(label_lines),
            retry_hint="will_retry" if pdf_path else "no_pdf",
        )
    return None


def _calendar_date(day: int, month: int, year: int):
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _parse_printed_date_token(day_s: str, month_s: str, year_s: str):
    day = int(day_s)
    month = int(month_s)
    year = int(year_s)
    if len(year_s) == 2:
        year = 2000 + year if year < 50 else 1900 + year
    if year < 1990 or year > 2099:
        return None
    return _calendar_date(day, month, year)


def extract_invoice_date(text: str):
    """Printed invoice date from the GST header. Source folder year is ignored."""
    text = normalize_ocr_text(text)
    labeled = re.search(
        r"Invoice\s*No\.?\s*(?:&\s*Date)?[\s:\-]*((?:(?!Invoice).){0,200})",
        text,
        flags=re.I | re.S,
    )
    windows = [labeled.group(1)] if labeled else []
    windows.append(text)

    for window in windows:
        after_number = re.search(
            r"\b20\d{9}\b\s*[-–:,]?\s*(\d{1,2}[./-]\d{1,2}[./-](?:\d{4}|\d{2}))",
            window,
        )
        candidates = []
        if after_number:
            candidates.append(after_number.group(1))
        candidates.extend(match.group(0) for match in _PRINTED_DATE_RE.finditer(window))
        for token in candidates:
            parsed = _PRINTED_DATE_RE.search(token)
            if not parsed:
                continue
            value = _parse_printed_date_token(*parsed.groups())
            if value:
                return value
    return None


_OFFICIAL_CUSTOMERS = None
_TOKEN_DROP = frozenset({"PVT", "LTD", "LIMITED", "PRIVATE", "LLC", "LLP", "CO"})
_TOKEN_FOLD = {
    "TECHNOLOGIES": "TECH",
    "TECHNOLOGY": "TECH",
    "ENGINEERING": "ENGG",
    "ENGINEERS": "ENGG",
    "MANUFACTURING": "MFG",
}


def load_official_customers():
    """Billed-to names from the mapping workbook, then customers.txt."""
    global _OFFICIAL_CUSTOMERS
    if _OFFICIAL_CUSTOMERS is None:
        from core.customer_master import load_customer_master

        master = load_customer_master()
        if master is not None:
            names = [item.official_name for item in master.customers]
        else:
            path = app_root() / "customers.txt"
            names = []
            if path.exists():
                for line in path.read_text(encoding="utf-8").splitlines():
                    name = " ".join(line.replace("\xa0", " ").split()).strip()
                    if name and not name.startswith("#"):
                        names.append(name)
        _OFFICIAL_CUSTOMERS = names
    return _OFFICIAL_CUSTOMERS


def customer_tokens(name: str):
    tokens = []
    for word in re.findall(r"[A-Za-z0-9]+", (name or "").upper()):
        word = _TOKEN_FOLD.get(word, word)
        if word and word not in _TOKEN_DROP:
            tokens.append(word)
    return tokens


def _tokens_in_order(needle, haystack):
    index = 0
    for token in haystack:
        if index < len(needle) and token == needle[index]:
            index += 1
    return index == len(needle)


def match_official_customer(text: str):
    """Map billed-to OCR text onto the official customer list."""
    text = normalize_ocr_text(text)
    billed = re.search(
        r"(?:Details\s+Of\s+Recipient|Billed\s+to)(.*?)(?:Consignee|GSTIN|Place\s+of\s+Supply|Invoice\s*No|$)",
        text,
        flags=re.I | re.S,
    )
    regions = []
    if billed:
        regions.append(billed.group(1))
    regions.append(text)

    customers = load_official_customers()
    for region in regions:
        hay = customer_tokens(region)
        compact = "".join(hay)
        best = None
        for name in customers:
            needle = customer_tokens(name)
            if not needle:
                continue
            if len(needle) == 1:
                matched = needle[0] in hay
            else:
                matched = _tokens_in_order(needle, hay) or ("".join(needle) in compact)
            if not matched:
                continue
            score = (len(needle), len(name))
            if best is None or score > best[0]:
                best = (score, name)
        if best:
            return best[1]
    return None


_COMPANY_RE = re.compile(
    r"([A-Z0-9][A-Z0-9 .&'()/-]*?(?:PRIVATE\s+LIMITED|PVT\.?\s*LTD\.?|LIMITED|LLP|LTD\.?))",
    re.I,
)
_LETTERHEAD_RE = re.compile(
    r"@|\b(?:EMAIL|PHONE|P\s*A\s*N|PAN|CIN|IRN|GAT\s*NO|VILLAGE)\b",
    re.I,
)


def _anchor_text(value: str) -> str:
    text = (value or "").upper().replace("0", "O").replace("1", "I")
    text = re.sub(r"[^A-Z]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _anchor_rank(value: str) -> int:
    """Higher is a better billed-to label. 0 means this line is not a label."""
    norm = _anchor_text(value)
    if not norm:
        return 0
    compact = norm.replace(" ", "")
    if "DETAIL" in compact and "RECIPIENT" in compact:
        return 3
    if "BILLEDTO" in compact:
        return 2
    if "BILLTO" in compact:
        return 1
    for phrase, rank in (("DETAILS OF RECIPIENT", 3), ("BILLED TO", 2), ("BILL TO", 1)):
        if len(norm) <= len(phrase) + 16 and SequenceMatcher(None, norm, phrase).ratio() >= 0.72:
            return rank
    return 0


def _is_section_stop(value: str) -> bool:
    norm = _anchor_text(value)
    compact = norm.replace(" ", "")
    if "CONSIGNEE" in compact:
        return True
    if "GSTIN" in compact and "UNIQUE" in compact:
        return True
    if "STATE" in compact and "CODE" in compact:
        return True
    return False


def _company_name(value: str) -> str:
    """Company line inside the billed-to section, or empty when it is not one."""
    text = re.sub(r"(?i)details\s+of\s+recipient|billed\s+to|bill\s+to", " ", value or "")
    text = re.sub(r"\s+", " ", text).strip(" ,:;-()")
    if not text or _LETTERHEAD_RE.search(text) or re.match(r"^\d", text):
        return ""
    if re.search(r"HIGHTEMP\s+FURNACES", text, re.I):
        return ""
    match = _COMPANY_RE.search(text)
    if not match:
        return ""
    return match.group(1).strip(" ,:;-.")


def _x_overlap(anchor: tuple[float, float, float, float], box: tuple[float, float, float, float]) -> float:
    overlap = min(anchor[2], box[2]) - max(anchor[0], box[0])
    if overlap <= 0:
        return 0.0
    width = max(1.0, anchor[2] - anchor[0])
    return overlap / width


def _same_row(anchor: tuple[float, float, float, float], box: tuple[float, float, float, float]) -> bool:
    top = max(anchor[1], box[1])
    bottom = min(anchor[3], box[3])
    height = max(1.0, min(anchor[3] - anchor[1], box[3] - box[1]))
    return bottom - top >= 0.4 * height


def customer_from_lines(lines: list[OcrLine]) -> str:
    """Company name in the billed-to column. Seller text above the label is ignored."""
    ranked = [( _anchor_rank(line.text), index, line) for index, line in enumerate(lines)]
    anchors = [item for item in ranked if item[0] > 0]
    if not anchors:
        return ""
    _rank, _index, anchor = max(anchors, key=lambda item: (item[0], -item[2].box[1]))
    on_label = _company_name(anchor.text)
    if on_label:
        return on_label

    ax0, ay0, ax1, ay1 = anchor.box
    stops = []
    for line in lines:
        if line is anchor or not _is_section_stop(line.text):
            continue
        if line.box[1] <= ay1:
            continue
        if _x_overlap(anchor.box, line.box) < 0.25 and not (line.box[0] >= ax0 - 8 and line.box[0] <= ax1 + 8):
            continue
        stops.append(line.box[1])
    stop_y = min(stops) if stops else float("inf")

    candidates = []
    for line in lines:
        if line is anchor:
            continue
        x0, y0, x1, y1 = line.box
        if y0 >= stop_y:
            continue
        in_column = _x_overlap(anchor.box, line.box) >= 0.25 or (ax0 - 8 <= x0 <= ax1)
        below = y0 >= ay0 - 2 and in_column
        # Name printed on the label's row, still in the billed-to column.
        beside = _same_row(anchor.box, line.box) and ax0 <= x0 <= ax1 + max(40.0, ax1 - ax0)
        if not below and not beside:
            continue
        name = _company_name(line.text)
        if name:
            candidates.append((y0, x0, name))
    if not candidates:
        return ""
    candidates.sort()
    return candidates[0][2]


def billed_to_region(text: str, lines: list[OcrLine] | None = None) -> str:
    """Customer company line from the billed-to section, not the seller letterhead."""
    if lines is None:
        lines = lines_from_plain_text(text)
    return customer_from_lines(lines)


def extract_customer_name(text: str):
    """Billed-to customer, using the official list when OCR matches it."""
    official = match_official_customer(text)
    if official:
        return official

    text = normalize_ocr_text(text)
    billed = re.search(
        r"(?:Details\s+Of\s+Recipient|Billed\s+to)(.*?)(?:Consignee|GSTIN|Place\s+of\s+Supply|Invoice\s*No|$)",
        text,
        flags=re.I | re.S,
    )
    regions = [billed.group(1)] if billed else []
    regions.append(text)

    company = re.compile(
        r"([A-Z][A-Z0-9 .&'/-]*?(?:PRIVATE\s+LIMITED|PVT\.?\s*LTD\.?|LTD\.?|LIMITED|LLP)\.?)",
        re.I,
    )
    for region in regions:
        for match in company.finditer(region):
            name = canonicalize_customer_name(match.group(1))
            if _usable_customer_name(name) and not re.search(r"HIGHTEMP", name, re.I):
                return name

    match = re.search(r"(PORITE\s*INDIA\s*PVT\.?\s*LTD\.?)", text, flags=re.I)
    if match:
        return canonicalize_customer_name("PORITE INDIA PVT.LTD.")
    return None


def canonicalize_customer_name(name: str) -> str:
    """One legal name: drop plant names, OCR junk, and billed/shipped duplicates."""
    name = _clean_customer_name(name).replace("]", ")")
    legal = re.search(
        r"(.+?(?:PRIVATE\s+LIMITED|PVT\.?\s*LTD\.?|LTD\.?|LIMITED|LLP)\.?)",
        name,
        flags=re.I,
    )
    if legal:
        name = legal.group(1)
    return _clean_customer_name(name)


def _clean_customer_name(name: str) -> str:
    return re.sub(r"\s+", " ", name or "").strip(" ,-")


def _usable_customer_name(name: str) -> bool:
    if not name or len(name) < 4:
        return False
    if re.fullmatch(r"\d+", name):
        return False
    if re.match(r"^\d+\s*,", name):
        return False
    if re.match(r"^(GSTIN|INVOICE|TAX|FORM|FILE\s+COPY)\b", name, re.I):
        return False
    return True


def looks_like_invoice_page(text: str):
    text = normalize_ocr_text(text)
    invoice_no = extract_invoice_number(text)
    gst_form = re.search(r"FORM\s+GST\s+INV", text, re.I) is not None
    tax_invoice = re.search(r"TAX\s+INVOICE|\bFILE\s+COPY\b", text, re.I) is not None
    delivery_only = (
        re.search(r"DELIVERY\s*CHALLAN|Original\s+For\s+Consignee", text, re.I)
        and not gst_form
        and not tax_invoice
    )
    if delivery_only:
        return False
    if invoice_no:
        return True
    return bool(gst_form and tax_invoice)


def date_folder_name_for(path: Path, root: Path) -> str:
    for parent in [path, *path.parents]:
        if DATE_FOLDER_RE.match(parent.name):
            return parent.name
        if parent == root:
            break
    return path.parent.parent.name


def process_invoice_file(source_pdf: Path, root: Path, output_root: Path):
    """Identify an invoice from page 1 and copy the COMPLETE source PDF when it is safe."""
    from core.pipeline import run_one

    return run_one(source_pdf, root, output_root)


def is_scan_date_folder(path: Path) -> bool:
    """A scan batch folder: ``01-Sep-26`` or any folder that directly contains Invoice/."""
    if not path.is_dir() or path.name.lower() == "invoice" or path.name.endswith("_done"):
        return False
    if any(parent.name.endswith("_done") for parent in path.parents):
        return False
    if DATE_FOLDER_RE.match(path.name):
        return True
    return invoice_child(path) is not None


def find_date_folders(root: Path):
    root = Path(root)
    found = []
    if root.is_dir() and is_scan_date_folder(root):
        found.append(root)
    if root.is_dir():
        for path in root.rglob("*"):
            if path.is_dir() and is_scan_date_folder(path):
                found.append(path)
    found_set = set(found)
    chosen = [path for path in found if not any(parent in found_set for parent in path.parents)]
    return sorted(
        chosen,
        key=lambda path: (parse_date_folder(path.name) or datetime.min.date(), str(path)),
    )


def invoice_child(date_folder: Path):
    if not date_folder.is_dir():
        return None
    try:
        children = list(date_folder.iterdir())
    except OSError:
        return None
    for child in children:
        if child.is_dir() and child.name.lower() == "invoice":
            return child
    return None


def invoice_pdfs_in(date_folder: Path):
    invoice_dir = None
    for child in date_folder.iterdir():
        if child.is_dir() and child.name.lower() == "invoice":
            invoice_dir = child
            break
    if invoice_dir is None:
        return None
    return sorted(
        p for p in invoice_dir.rglob("*")
        if p.is_file() and p.suffix.lower() == ".pdf" and not p.name.startswith(".")
    )


def extract_zip(archive: Path, dest: Path) -> Path:
    dest.mkdir(parents=True, exist_ok=True)
    dest_resolved = dest.resolve()
    with zipfile.ZipFile(archive) as zf:
        for info in zf.infolist():
            target = (dest / info.filename).resolve()
            if dest_resolved not in target.parents and target != dest_resolved:
                raise ValueError(f"Unsafe zip entry: {info.filename}")
        zf.extractall(dest)
    return dest


_ZIP_SOURCE_MARKER = ".invoice_sorter_zip_source"


def _write_zip_source_marker(dest: Path, archive: Path) -> None:
    stat = archive.stat()
    (dest / _ZIP_SOURCE_MARKER).write_text(
        f"{stat.st_mtime}\n{stat.st_size}\n",
        encoding="utf-8",
    )


def _zip_extract_stale(archive: Path, dest: Path) -> bool:
    """True when the zip must be extracted again (missing tree or archive changed)."""
    if not dest.exists() or not find_date_folders(dest):
        return True
    try:
        zip_stat = archive.stat()
    except OSError:
        return True
    marker = dest / _ZIP_SOURCE_MARKER
    if marker.is_file():
        try:
            mtime_s, size_s = marker.read_text(encoding="utf-8").splitlines()[:2]
            if float(mtime_s) == zip_stat.st_mtime and int(size_s) == zip_stat.st_size:
                return False
        except (ValueError, OSError):
            pass
        return True
    try:
        return zip_stat.st_mtime > dest.stat().st_mtime
    except OSError:
        return True


def resolve_input(path: Path) -> Path:
    path = path.expanduser()
    if path.is_file() and path.suffix.lower() == ".zip":
        dest = path.parent / f"{path.stem}_extracted"
        if _zip_extract_stale(path, dest):
            if dest.exists():
                shutil.rmtree(dest)
            extract_zip(path, dest)
            _write_zip_source_marker(dest, path)
        return dest
    if path.is_dir():
        return path
    raise FileNotFoundError(path)


EXCEPTION_REPORT_NAME = "invoice_sorter_exceptions.xlsx"
_REPORT_COLUMNS = (
    "status",
    "source_file",
    "date_folder",
    "invoice_number",
    "customer",
    "year",
    "source_pages",
    "reason",
)


def _report_rows(results, statuses):
    return [row for row in results if row.get("status") in statuses]


def write_exception_report(results, dest) -> Path | None:
    """Excel of files that could not be read (REVIEW) or were skipped.

    dest may be a Path or a binary file object. Always writes Summary,
    Could_not_read, and Skipped sheets so the user can download a report
    even when every invoice copied successfully.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    copied = sum(r.get("status") == "COPIED" for r in results)
    review = _report_rows(results, {"REVIEW", "FAILED"})
    skipped = _report_rows(results, {"SKIPPED"})

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Metric", "Count"])
    summary.append(["Copied", copied])
    summary.append(["Could not read (REVIEW)", len(review)])
    summary.append(["Skipped", len(skipped)])
    summary.append(["Total rows", len(results)])
    summary["A1"].font = Font(bold=True)
    summary["B1"].font = Font(bold=True)

    def _fill(sheet_name, rows):
        sheet = wb.create_sheet(sheet_name)
        headers = [
            "Status",
            "Source file",
            "Source day folder",
            "Invoice number",
            "Customer",
            "Year",
            "Source pages",
            "Reason",
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([row.get(key, "") or "" for key in _REPORT_COLUMNS])
        for index, header in enumerate(headers, start=1):
            extra = max((len(str(row.get(_REPORT_COLUMNS[index - 1], "") or "")) for row in rows), default=0)
            sheet.column_dimensions[get_column_letter(index)].width = min(60, max(len(header) + 2, extra + 2))

    _fill("Could_not_read", review)
    _fill("Skipped", skipped)

    wb.save(dest)
    return dest if isinstance(dest, Path) else None


def exception_report_bytes(results) -> bytes:
    buffer = io.BytesIO()
    write_exception_report(results, buffer)
    return buffer.getvalue()


def process(root: Path, output_root: Path, progress=None):
    """Analyze the batch, copy safe matches, and mark finished source folders _done."""
    from core.pipeline import run_batch

    return run_batch(root, output_root, progress=progress, execute=True)


def zip_output_tree(output_root: Path) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for file in sorted(output_root.rglob("*")):
            if file.is_file():
                zf.write(file, file.relative_to(output_root).as_posix())
    return buffer.getvalue()


def process_uploaded_zip(zip_bytes: bytes, filename: str, work_dir: Path, progress=None):
    safe_zip_name = Path(filename or "invoices.zip").name
    if not safe_zip_name.lower().endswith(".zip"):
        safe_zip_name += ".zip"
    work_dir.mkdir(parents=True, exist_ok=True)
    source = work_dir / safe_zip_name
    source.write_bytes(zip_bytes)
    output_root = work_dir / "sorted"
    output_root.mkdir(parents=True, exist_ok=True)
    results = process(source, output_root, progress=progress)
    return results, zip_output_tree(output_root)
